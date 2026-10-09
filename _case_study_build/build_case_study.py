from pathlib import Path
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(__file__).resolve().parent.parent
BUILD = ROOT / '_case_study_build'
OUT = ROOT / 'output' / 'case_study'
OUT.mkdir(parents=True, exist_ok=True)

PAPER1 = 'https://doi.org/10.1017/wet.2024.70'
PAPER2 = 'https://doi.org/10.1002/agg2.70055'
PAPER3 = 'https://doi.org/10.1002/agg2.70324'
TECH = 'https://www.deere.com/en-us/products-solutions/sprayers-and-applicators/see-spray'
PHOTO1 = 'https://newsroom.aem.org/john-deere-at-ag-on-the-mall-2024/'
PHOTO2 = 'https://www.deere.com/en/technology-products/precision-ag-technology/precision-upgrades/sprayer-upgrades/getting-started-ultimate/'

def font(size, bold=False):
    path = 'C:/Windows/Fonts/arialbd.ttf' if bold else 'C:/Windows/Fonts/arial.ttf'
    return ImageFont.truetype(path, size)

def diagram():
    im = Image.new('RGB', (2200, 620), '#ffffff')
    d = ImageDraw.Draw(im)
    green, blue, ink = '#236341', '#245b83', '#182b30'
    boxes = [
        (50, 115, 480, 420, '01  SENSE', 'Boom cameras +\nheight sensors', 'Capture crop and weed images'),
        (610, 115, 1040, 420, '02  DECIDE', 'Onboard AI vision', 'Identify weeds and locate targets'),
        (1170, 115, 1600, 420, '03  ACT', 'Nozzle controller', 'Spray selected spots only'),
        (1730, 115, 2160, 420, '04  RECORD', 'Application map', 'Review coverage and settings'),
    ]
    for i, (x1,y1,x2,y2,hd,main,sub) in enumerate(boxes):
        color = green if i in (0,2) else blue
        d.rounded_rectangle((x1,y1,x2,y2), radius=28, fill='#f7faf8', outline=color, width=6)
        d.rounded_rectangle((x1,y1,x2,y1+78), radius=26, fill=color)
        d.rectangle((x1,y1+45,x2,y1+78), fill=color)
        d.text((x1+25,y1+20), hd, font=font(34,True), fill='white')
        yy = y1+112
        for line in main.split('\n'):
            d.text((x1+25,yy), line, font=font(41,True), fill=ink)
            yy += 48
        d.text((x1+25,y2-73), sub, font=font(24), fill='#445b5d')
    for x in (505,1065,1625):
        d.line((x,265,x+78,265), fill=ink, width=9)
        d.polygon([(x+78,265),(x+48,246),(x+48,284)],fill=ink)
    d.text((60,482), 'Operator sets crop, sensitivity and spray program; field conditions affect the result.', font=font(31), fill=ink)
    d.text((60,537), 'Architecture synthesized from the product description and field studies [1–4].', font=font(25), fill='#53696b')
    p=BUILD/'architecture.png'
    im.save(p)
    return p

def hyperlink(paragraph, label, url, size=8):
    part=paragraph.part
    rid=part.relate_to(url, RT.HYPERLINK, is_external=True)
    h=OxmlElement('w:hyperlink'); h.set(qn('r:id'),rid)
    r=OxmlElement('w:r')
    pr=OxmlElement('w:rPr')
    col=OxmlElement('w:color'); col.set(qn('w:val'),'245B83'); pr.append(col)
    ul=OxmlElement('w:u'); ul.set(qn('w:val'),'single'); pr.append(ul)
    sz=OxmlElement('w:sz'); sz.set(qn('w:val'),str(int(size*2))); pr.append(sz)
    r.append(pr)
    t=OxmlElement('w:t'); t.text=label; r.append(t); h.append(r)
    paragraph._p.append(h)

def add_body(doc, heading, text):
    doc.add_heading(heading, 2)
    p=doc.add_paragraph(text)
    p.style='Normal'
    return p

doc=Document()
sec=doc.sections[0]
sec.page_width=Inches(8.5); sec.page_height=Inches(11)
sec.top_margin=Inches(.62); sec.bottom_margin=Inches(.55)
sec.left_margin=Inches(.72); sec.right_margin=Inches(.72)
sec.header_distance=Inches(.25); sec.footer_distance=Inches(.28)

styles=doc.styles
styles['Normal'].font.name='Aptos'; styles['Normal'].font.size=Pt(9.5); styles['Normal'].font.color.rgb=RGBColor(28,39,42)
styles['Normal'].paragraph_format.space_after=Pt(4)
styles['Normal'].paragraph_format.line_spacing=1.08
for name,size in [('Title',19),('Heading 1',12),('Heading 2',10)]:
    s=styles[name]; s.font.name='Aptos Display' if name=='Title' else 'Aptos'; s.font.size=Pt(size)
    s.font.bold=True; s.font.color.rgb=RGBColor(0,0,0)
    s.paragraph_format.space_before=Pt(8 if name=='Heading 2' else 0)
    s.paragraph_format.space_after=Pt(2)

head=sec.header.paragraphs[0]
head.text='ARTIFICIAL INTELLIGENCE & SUSTAINABLE DEVELOPMENT  |  TA 1 CASE STUDY'
head.style='Normal'; head.runs[0].font.size=Pt(7.5); head.runs[0].font.bold=True
head.runs[0].font.color.rgb=RGBColor(64,93,96)
footer=sec.footer.paragraphs[0]
footer.text='AI for Agriculture and Environment  •  Research links and image credits on page 4'
footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
footer.runs[0].font.size=Pt(7)

doc.add_paragraph('AI Guided Precision Spraying in Soybean Fields', 'Title')
p=doc.add_paragraph('A real-world case study of John Deere See & Spray Ultimate')
p.runs[0].font.size=Pt(10); p.runs[0].font.color.rgb=RGBColor(35,99,65)
p=doc.add_paragraph('Theme: AI for Agriculture and Environment  |  Relevant goals: SDG 2 (sustainable agriculture) and SDG 12 (responsible production)')
p.runs[0].font.size=Pt(8.5)

doc.add_heading('Case overview', 1)
add_body(doc,'The central question',
    'Can a sprayer use artificial intelligence to identify weeds in real time and apply herbicide only where it is needed, while preserving effective weed control? John Deere’s See & Spray Ultimate provides a practical case because it has been deployed on farms and evaluated in published field studies. Its performance can be compared with conventional broadcast spraying, which treats the whole field. [1–4]')
add_body(doc,'Case profile',
    'Organisation and application: John Deere / See & Spray Ultimate, an AI-assisted agricultural sprayer. Users: farmers and spray operators. Crop examined here: soybean. Evidence base: replicated field experiments in Arkansas, Mississippi and Nebraska, plus anonymised U.S. commercial application records. Core input: live field images. Core output: decisions to activate selected spray nozzles. This combination makes the case technically clear and allows the claimed benefits to be checked against measured results. [1–4]')
add_body(doc,'1. Background / Problem',
    'Weeds compete with crops for sunlight, water and nutrients. Farmers therefore apply herbicides to protect yield. In a broadcast program, a spray boom applies product across the entire treated width, including crop rows and areas without visible weeds. When weeds are clustered in patches, some of this application is unnecessary. Chemical inputs have a cost, and reducing avoidable use can support more efficient resource management. The central difficulty is that a machine moving through a field must detect small weeds quickly enough to control them without missing targets. [1, 2]')
add_body(doc,'Why this is a sustainability problem',
    'A useful AI system must improve selectivity, not merely automate spraying. Less targeted herbicide applied means less purchased chemical enters the production process. It may reduce the potential for off-target exposure, but actual changes in water quality, soil health or human exposure require separate measurement. This case therefore evaluates the directly observed outcomes—sprayed area, herbicide used, weed control and yield—before making environmental claims. [1–3]')
add_body(doc,'2. AI Based Solution',
    'See & Spray combines cameras, sensors, onboard image processing and individually controlled nozzles. Images of plants and soil are analysed by a trained machine-vision model to distinguish weeds from the crop or background. A detection is converted into a spray decision at the relevant location, and the controller activates selected nozzles as the boom passes. The operator chooses the crop, application program and detection settings, then checks the field results. The machine performs rapid recognition and actuation; agronomic decisions and supervision remain human responsibilities. [1, 4]')
add_body(doc,'What makes the solution “AI”',
    'The important step is visual classification: the system has to infer from image patterns which plants should be treated. A conventional timer or fixed GPS rule cannot make that distinction from a new camera image. The AI output is actionable only when connected to sensing, location, nozzle timing and a usable operator interface. [1, 4]')

doc.add_page_break()
doc.add_heading('System architecture and equipment', 1)
p=doc.add_paragraph('Figure 1. Functional architecture of targeted spraying (author-created synthesis from [1–4]).')
p.runs[0].font.size=Pt(8.5)
doc.add_picture(str(diagram()), width=Inches(7.05))
add_body(doc,'Step-by-step operation',
    'First, cameras on the boom observe the ground ahead of the spray nozzles. Height sensors help maintain the operating position. Second, the onboard vision system identifies likely weeds and their positions. Third, the controller matches each target with the appropriate nozzle and spray timing as the vehicle moves. Fourth, the operator reviews application records and adjusts the program when field conditions change. A targeted pass may still involve a separate broadcast residual herbicide; “spot spray” does not mean every chemical on the farm is applied only to weeds. [1, 4]')
doc.add_heading('Real equipment photographs', 1)
t=doc.add_table(rows=1, cols=2)
t.autofit=False; t.alignment=WD_TABLE_ALIGNMENT.CENTER
t.columns[0].width=Inches(3.42); t.columns[1].width=Inches(3.42)
for i,(filename,caption,url) in enumerate([
    ('sprayer.jpg','Figure 2. See & Spray Ultimate operating in a crop field. Photo: AEM / John Deere.',PHOTO1),
    ('nozzle.jpg','Figure 3. John Deere boom nozzle assembly. Photo: John Deere.',PHOTO2)
]):
    c=t.cell(0,i); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.TOP
    c.paragraphs[0].add_run().add_picture(str(BUILD/filename),width=Inches(3.36),height=Inches(1.80))
    p=c.add_paragraph(caption); p.paragraph_format.space_after=Pt(0)
    for r in p.runs: r.font.size=Pt(7.5)
    q=c.add_paragraph(); hyperlink(q,'Image source',url,7.5)
p=doc.add_paragraph('Photos show commercial equipment, not the cited research plots. The architecture is a functional explanation, not a manufacturer circuit diagram.')
p.runs[0].font.size=Pt(7.5)

doc.add_page_break()
doc.add_heading('3. Application / Implementation', 1)
add_body(doc,'Arkansas and Mississippi: field comparison',
    'Researchers conducted soybean experiments in 2021–22 at Keiser, Arkansas, and Greenville, Mississippi. They compared traditional broadcast herbicide programs with targeted applications using See & Spray. The treatments included preemergence herbicide followed by early and mid-postemergence applications. Some programs broadcast a residual herbicide while targeting the postemergence product; others used targeted applications in a single-tank approach. The study assessed the treated area, weed control, soybean injury and yield. This design matters because a fair comparison must measure crop protection as well as chemical savings. [1]')
add_body(doc,'Nebraska: a contrasting soybean field',
    'A separate 2023 experiment near Mead, Nebraska used a John Deere 612R See & Spray Ultimate in soybean. The reported machine had 36 cameras and 96 nozzles across a 36.6 m boom. Researchers compared spot-spray and broadcast treatments in a field with substantial weed pressure. Here the technology still controlled weeds comparably in most treatments, but there were few weed-free areas to skip. This is an important real-world stress case for the savings claim. [2]')
add_body(doc,'Commercial deployment',
    'A later analysis studied 11,140 anonymised commercial sprayer application observations from the United States. The records included weed area, sprayed area, crop, field size, platform, date and region. They show that the technology has moved beyond small experimental plots. However, the dataset did not include every machine setting or nozzle choice, so it cannot explain every individual result. [3]')
add_body(doc,'What implementation requires on a farm',
    'Before spraying, an operator must select an appropriate crop and herbicide program, confirm equipment readiness and choose detection settings. During operation, camera visibility, boom height, speed and nozzle performance affect whether the intended weed receives adequate coverage. Afterwards, the application record and field scouting help determine whether another treatment is needed. The field studies therefore evaluated the full agricultural outcome, not image classification accuracy alone. A detection system that saves chemical but leaves uncontrolled weeds would not solve the original problem. [1–4]')
add_body(doc,'How the studies were interpreted',
    'The percentage saved is measured against a broadcast application under the study’s treatment design. A lower sprayed area is meaningful only alongside acceptable weed control and crop response. The Arkansas–Mississippi study reports several application timings; the Nebraska work used a different field, season and weed pressure. The commercial observations combine many farms and crops. They are complementary pieces of evidence, not repetitions of one identical experiment. [1–3]')

doc.add_heading('Field evidence: what was actually measured', 1)
t=doc.add_table(rows=4,cols=3); t.style='Table Grid'
rows=[
    ('Setting','Observed result','Interpretation'),
    ('AR + MS soybean trials, 2021–22 [1]','28.4%–62.4% less targeted postemergence herbicide; ≥93% reported weed control','Strong savings were possible with comparable control.'),
    ('NE soybean trial, 2023 [2]','1.1%–5.4% early non-residual savings; 0% late savings','Dense weeds left little area unsprayed.'),
    ('U.S. commercial records, 2023–25 [3]','11,140 observations; 58.4% less area sprayed on average','Broad use, but results vary by field and settings.'),
]
for row,values in zip(t.rows,rows):
    for cell,value in zip(row.cells,values):
        cell.text=value
        for p in cell.paragraphs:
            for r in p.runs: r.font.size=Pt(8)
for c in t.rows[0].cells:
    for p in c.paragraphs:
        for r in p.runs: r.font.bold=True
p=doc.add_paragraph('The percentages refer to different measures and study settings; they should not be added or treated as one guaranteed saving for every farm.')
p.runs[0].font.size=Pt(8); p.runs[0].font.italic=True

doc.add_page_break()
doc.add_heading('4. Benefits / Impact', 1)
add_body(doc,'Measured benefits',
    'The main demonstrated benefit is a reduction in targeted herbicide or sprayed area when weeds are sparse enough for the sprayer to leave parts of a field untreated. In the Arkansas–Mississippi trials, reported control for several weed species remained at least 93%, with no yield increase attributable to targeted application. The commercial analysis found an average 58.4% reduction in area sprayed across its recorded applications. Thus the best-supported claim is improved input efficiency while maintaining comparable weed control in suitable conditions. [1, 3]')
add_body(doc,'Sustainable-development relevance',
    'The case supports SDG 2 by exploring a more efficient way to protect crop production, and SDG 12 by using an agricultural input more selectively. Reduced herbicide applied could reduce potential environmental loading and operating expense. The cited studies did not directly measure long-term soil health, water contamination, biodiversity or a general yield gain; those outcomes should be treated as possible downstream effects to investigate, not proven impacts of this case. [1–3]')
doc.add_heading('5. Challenges / Limitations', 1)
add_body(doc,'Field and technical limits',
    'High weed density is the clearest limit. In the Nebraska soybean field, early non-residual savings were only 1.1%–5.4%, and late applications saved none. Detection may also be affected by small or obscured weeds, crop canopy and operating conditions. Spray coverage can require adjacent nozzles or overlap, which increases treated area. Detection sensitivity, buffer choices and nozzle fan angle change the trade-off between saving chemical and covering weeds. [2, 3]')
add_body(doc,'Economic and management limits',
    'The sprayer requires capital, maintenance, calibration and operator training. A 2026 economic analysis estimated that, under its assumptions, herbicide spending per application had to exceed about US$27 per hectare for See & Spray Premium or US$39 per hectare for Ultimate to break even against broadcast spraying. These are study-specific thresholds, not universal prices. Farmers also need an appropriate herbicide program and must monitor for escapes rather than accepting a lower application total as success by itself. [3]')
add_body(doc,'Practical recommendation',
    'A sensible adoption decision would begin with a representative field trial: record weed density, targeted area sprayed, herbicide cost, weed-control score and crop yield alongside a comparable broadcast treatment. Review the result across more than one application timing. If savings disappear as weed pressure rises, improve preventive weed management before expecting the AI sprayer to deliver a large reduction. This recommendation follows from the contrast between the field studies and the commercial analysis. [1–3]')
doc.add_heading('6. Conclusion', 1)
p=doc.add_paragraph('See & Spray is a useful example of AI for sustainable agriculture because a visual prediction is connected to a physical action in an operating farm machine. The strongest evidence shows that it can avoid substantial targeted spraying when weeds are patchy while retaining comparable control. The contrasting Nebraska result prevents an exaggerated claim: AI cannot create large savings in a field that is already densely infested. Successful use depends on sound weed management, suitable settings, field monitoring and an economic case for the farm. [1–3]')

doc.add_heading('Linked references and photograph credits', 1)
refs=[
    ('[1] Avent et al. (2024). Comparing herbicide application methods with See & Spray technology in soybean. Weed Technology. DOI 10.1017/wet.2024.70',PAPER1),
    ('[2] Leise et al. (2025). Evaluating precision sprayers for targeted weed management in corn and soybean. Agrosystems, Geosciences & Environment. DOI 10.1002/agg2.70055',PAPER2),
    ('[3] Avent et al. (2026). Targeted herbicide cost requirements for adoption of See & Spray technology. Agrosystems, Geosciences & Environment. DOI 10.1002/agg2.70324',PAPER3),
    ('[4] John Deere. See & Spray technology description.',TECH),
    ('Photo for Figure 2: Association of Equipment Manufacturers / John Deere.',PHOTO1),
    ('Photo for Figure 3: John Deere sprayer upgrades page.',PHOTO2),
]
for label,url in refs:
    p=doc.add_paragraph(style='Normal'); p.paragraph_format.space_after=Pt(1)
    hyperlink(p,label,url,7.5)

doc.save(OUT/'AI_Precision_Spraying_Case_Study.docx')
print(OUT/'AI_Precision_Spraying_Case_Study.docx')
